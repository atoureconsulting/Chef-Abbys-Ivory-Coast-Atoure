"""Build 03-budget/ground-budget.xlsx (live formulas) and print a summary.

Edit LINES / ASSUMPTIONS here, or edit the spreadsheet directly — the sheet
recalculates on its own. Figures are pre-quote estimates in XOF.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "03-budget" / "ground-budget.xlsx"

ASSUMPTIONS = [  # (label, value, note)
    ("Tour party (people)", 6, "Chef Abbys + manager + 2 camera + photo + assistant — CONFIRM"),
    ("Hotel nights", 4, "Nights of 18–21 Oct (arrive the evening before); depart 22 Oct to Monrovia"),
    ("Programme days", 3, "19–21 Oct"),
    ("Vehicle days", 5, "18 Oct airport pick-up through 22 Oct airport drop"),
    ("XOF per USD", 565, "Check on the day"),
    ("XOF per EUR", 655.957, "Fixed peg"),
    ("Contingency", 0.10, "On ground costs"),
    ("AToure curation & management fee", 0.15, "PLACEHOLDER — Baba to set"),
]
# Named cells: B3.. in order
A = {k: f"Assumptions!$B${i+3}" for i, (k, _, _) in enumerate(ASSUMPTIONS)}
PAX, NIGHTS, DAYS, VDAYS = A["Tour party (people)"], A["Hotel nights"], A["Programme days"], A["Vehicle days"]

# v0.2 — merged programme (team "Parcours" spine + Olivia Yacé goûter + Jiaan Wu maquis + ministries)
# (section, item, unit cost XOF, qty (number or formula string), tier, funding route, note)
LINES = [
    ("Arrival & stay", "Airport protocol welcome + fast-track (arrival & departure)", 200_000, 2, "Core", "In-kind: ministry protocol", ""),
    ("Arrival & stay", "L'Avenue Hôtel (or equivalent), single rooms", 150_000, f"={PAX}*{NIGHTS}", "Core", "In-kind: hotel partner", "Rate to quote"),
    ("Arrival & stay", "Per diems — meals outside programme", 25_000, f"={PAX}*{NIGHTS}", "Core", "Paid", ""),
    ("Arrival & stay", "Local SIM + data", 15_000, f"={PAX}", "Core", "In-kind: telco", "Orange / MTN / Moov"),
    ("Ground transport", "Talent SUV with driver, fuel incl.", 100_000, f"={VDAYS}", "Core", "In-kind: mobility partner", ""),
    ("Ground transport", "Crew minibus with driver, fuel incl.", 110_000, f"={VDAYS}", "Core", "In-kind: mobility partner", ""),
    ("Ground transport", "Fuel & tolls supplement — Grand-Bassam run", 40_000, 1, "Core", "Paid", ""),
    ("Ground team", "Bilingual fixer / on-camera interpreter (FR↔EN)", 80_000, f"={DAYS}", "Core", "Paid", "Essential"),
    ("Ground team", "Protocol & security officers (×2)", 50_000, f"=2*{DAYS}", "Core", "Paid", ""),
    ("Ground team", "Runner / production assistant", 25_000, f"={DAYS}", "Core", "Paid", ""),
    ("Day 1 · The lagoon", "Attiéké entrepreneur — workshop honorarium", 300_000, 1, "Core", "Paid", "Women-led brand, to identify"),
    ("Day 1 · The lagoon", "Lagoon beach location, set & access", 150_000, 1, "Core", "Paid", "Site to confirm"),
    ("Day 1 · The lagoon", "Garba ingredients & tasting", 60_000, 1, "Core", "Paid", ""),
    ("Day 1 · Snack time", "Goûter ingredients & kit (~40 children)", 250_000, 1, "Core", "Sponsor", "With Fondation Olivia Yacé"),
    ("Day 1 · Snack time", "Burners, gas, aprons, hygiene", 150_000, 1, "Core", "Paid", ""),
    ("Day 1 · Snack time", "Child-safeguarding officer + consent forms", 100_000, 1, "Core", "Paid", "Non-negotiable"),
    ("Day 1 · Maquis night", "Reserved space + live band", 350_000, 1, "Core", "Sponsor", ""),
    ("Day 1 · Maquis night", "Food & drinks (party + ~6 guests)", 15_000, f"={PAX}+6", "Core", "Sponsor", ""),
    ("Day 1 · Maquis night", "Jiaan Wu — co-host fee [TBC]", 300_000, 1, "Core", "Paid / exposure swap", "Placeholder until quoted"),
    ("Day 2 · Grand-Bassam", "Zeinab Bancé — foutou workshop fee", 500_000, 1, "Core", "Paid", "Placeholder until quoted"),
    ("Day 2 · Grand-Bassam", "Lagoon-side garden location", 200_000, 1, "Core", "In-kind: venue", ""),
    ("Day 2 · Grand-Bassam", "Ingredients, mortar, artisan crockery", 150_000, 1, "Core", "Paid", ""),
    ("Day 2 · Grand-Bassam", "Heritage-quarter guide", 50_000, 1, "Core", "In-kind: Ministry of Culture", ""),
    ("Day 3 · Cocoa & zaouli", "Maison Mandjou atelier (Ahoua Touré)", 400_000, 1, "Core", "In-kind / co-brand", ""),
    ("Day 3 · Cocoa & zaouli", "Bean-to-bar demo — partner chocolatier", 300_000, 1, "Core", "In-kind / co-brand", "e.g. Le Chocolatier Ivoirien"),
    ("Day 3 · Cocoa & zaouli", "Gouro zaouli troupe + musicians at KAVA", 600_000, 1, "Core", "In-kind: Ministry of Culture", ""),
    ("Day 3 · Cocoa & zaouli", "Customary gifts & drinks — troupe", 100_000, 1, "Core", "Paid", ""),
    ("Day 3 · Closing table", "Closing dinner (party + ~10 guests)", 50_000, f"={PAX}+10", "Core", "Sponsor", "Ministers, Olivia, chefs"),
    ("Day 3 · Closing table", "Courtesy gifts (ministries, foundation)", 100_000, 1, "Core", "Paid", ""),
    ("Production support", "Licensed drone operator + ANAC authorisation", 350_000, 2, "Core", "Paid", "Days 1 & 2"),
    ("Production support", "Welcome kit — pagne, AToure merch", 25_000, f"={PAX}", "Core", "Paid", ""),
    # Optional / upgrades
    ("Upgrades", "CASOF × Abidjan masterclass (20–30 young cooks)", 300_000, 1, "Optional", "Sponsor", ""),
    ("Upgrades", "Cocoa farm visit (half day)", 300_000, 1, "Optional", "In-kind: Conseil Café-Cacao", ""),
    ("Upgrades", "Sunset pinasse on the Ébrié lagoon", 400_000, 1, "Optional", "Sponsor", ""),
    ("Upgrades", "Local BTS photographer / videographer", 200_000, f"={DAYS}", "Optional", "Paid", "Sponsor & ministry assets"),
    ("Upgrades", "Legacy seed gift — Fondation Olivia Yacé", 1_000_000, 1, "Optional", "Sponsor", "Girls' school fees / goûter"),
    ("Travel (outside ground)", "Flights Accra → Abidjan", 200_000, f"={PAX}", "Travel", "In-kind: airline", "May be on tour budget already"),
    ("Travel (outside ground)", "Flights Abidjan → Monrovia (Air Côte d'Ivoire)", 280_000, f"={PAX}", "Travel", "In-kind: Air Côte d'Ivoire", "Check days for 22 Oct"),
    ("Travel (outside ground)", "E-visas for non-ECOWAS crew", 50_000, 2, "Travel", "Paid", "Ghanaian passports: none needed"),
]

def qty_value(q, a):
    if isinstance(q, (int, float)):
        return q
    expr = q.lstrip("=")
    for k, ref in A.items():
        expr = expr.replace(ref, repr(a[k]))
    return eval(expr)

GOLD, BLACK, CREAM, PALE = "C8A951", "0D0C0A", "FAF6EE", "F5EDD6"
hdr_font = Font(name="Jost", bold=True, color=CREAM)
hdr_fill = PatternFill("solid", fgColor=BLACK)
sec_fill = PatternFill("solid", fgColor=PALE)
thin = Side(style="thin", color="E5D9BE")
box = Border(bottom=thin)
money = '#,##0" XOF"'
usd = '"$"#,##0'

wb = Workbook()
ws_a = wb.active
ws_a.title = "Assumptions"
ws_a["A1"] = "Chef Abbys × Côte d'Ivoire — budget assumptions (edit blue cells)"
ws_a["A1"].font = Font(name="Cormorant Garamond", size=16, bold=True)
for c, h in zip("ABC", ["Assumption", "Value", "Note"]):
    ws_a[f"{c}2"] = h; ws_a[f"{c}2"].font = hdr_font; ws_a[f"{c}2"].fill = hdr_fill
for i, (k, v, n) in enumerate(ASSUMPTIONS):
    r = i + 3
    ws_a[f"A{r}"], ws_a[f"B{r}"], ws_a[f"C{r}"] = k, v, n
    ws_a[f"B{r}"].font = Font(color="1F4FBF", bold=True)
    if isinstance(v, float) and v < 1:
        ws_a[f"B{r}"].number_format = "0%"
ws_a.column_dimensions["A"].width = 36; ws_a.column_dimensions["B"].width = 12; ws_a.column_dimensions["C"].width = 62

ws = wb.create_sheet("Budget")
cols = ["Section", "Line item", "Unit cost (XOF)", "Qty", "Total (XOF)", "Total (USD)", "Tier", "Include? (Y/N)", "Funding route", "Note"]
widths = [22, 58, 16, 8, 16, 12, 10, 14, 32, 44]
ws["A1"] = "Chef Abbys × Côte d'Ivoire · 19–21 Oct 2026 · Ground budget v0.2 (pre-quote estimate)"
ws["A1"].font = Font(name="Cormorant Garamond", size=16, bold=True)
for i, (h, w) in enumerate(zip(cols, widths)):
    c = ws.cell(row=2, column=i + 1, value=h); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[chr(65 + i)].width = w
dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=False); ws.add_data_validation(dv)

a_vals = {k: v for k, v, _ in ASSUMPTIONS}
r = 3
first = r
totals = {"Core": 0, "Optional": 0, "Travel": 0}
for sec, item, unit, q, tier, route, note in LINES:
    inc = "Y" if tier == "Core" else "N"
    row = [sec, item, unit, q, f"=IF(H{r}=\"Y\",C{r}*D{r},0)", f"=E{r}/{A['XOF per USD']}", tier, inc, route, note]
    for j, v in enumerate(row):
        c = ws.cell(row=r, column=j + 1, value=v); c.border = box
    ws[f"C{r}"].number_format = money; ws[f"E{r}"].number_format = money; ws[f"F{r}"].number_format = usd
    ws[f"H{r}"].font = Font(color="1F4FBF", bold=True); dv.add(ws[f"H{r}"])
    if tier != "Core":
        for j in range(1, 11): ws.cell(row=r, column=j).fill = sec_fill
    totals[tier] += unit * qty_value(q, a_vals)
    r += 1
last = r - 1
r += 1
def tot(label, formula, bold=False):
    global r
    ws[f"B{r}"] = label; ws[f"E{r}"] = formula; ws[f"F{r}"] = f"=E{r}/{A['XOF per USD']}"
    ws[f"E{r}"].number_format = money; ws[f"F{r}"].number_format = usd
    if bold:
        for c in "BEF": ws[f"{c}{r}"].font = Font(bold=True, color=BLACK); ws[f"{c}{r}"].fill = PatternFill("solid", fgColor=GOLD)
    r += 1; return r - 1
g = tot("Ground costs — included lines (excl. travel)", f'=SUMIFS(E{first}:E{last},G{first}:G{last},"<>Travel")')
c_ = tot("Contingency", f"=E{g}*{A['Contingency']}")
f_ = tot("AToure curation & management fee", f"=(E{g}+E{c_})*{A['AToure curation & management fee']}")
t_ = tot("TOTAL GROUND PROPOSITION", f"=E{g}+E{c_}+E{f_}", bold=True)
r += 1
tot("Travel lines included (reference, outside ground total)", f'=SUMIFS(E{first}:E{last},G{first}:G{last},"Travel")')
ws.freeze_panes = "C3"

ws_s = wb.create_sheet("Scenarios")
ws_s["A1"] = "Quick scenarios (static snapshot at build time — Budget sheet is the live source)"
ws_s["A1"].font = Font(bold=True)
ct, cont, fee, fx = a_vals["Contingency"], a_vals["AToure curation & management fee"], a_vals["AToure curation & management fee"], a_vals["XOF per USD"]
def full(x): return x * (1 + ct) * (1 + fee)
rows = [("Scenario", "XOF", "USD"),
        ("Core programme", full(totals["Core"]), full(totals["Core"]) / fx),
        ("Core + all upgrades", full(totals["Core"] + totals["Optional"]), full(totals["Core"] + totals["Optional"]) / fx),
        ("Travel lines (if we must fund)", totals["Travel"], totals["Travel"] / fx)]
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = ws_s.cell(row=i + 2, column=j + 1, value=round(v) if isinstance(v, float) else v)
        if i == 0: c.font = hdr_font; c.fill = hdr_fill
        elif j == 1: c.number_format = money
        elif j == 2: c.number_format = usd
ws_s.column_dimensions["A"].width = 36; ws_s.column_dimensions["B"].width = 18; ws_s.column_dimensions["C"].width = 12

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
secs={}
for sec,item,unit,q,tier,_,_ in LINES:
    if tier=="Core": secs[sec.split(" · ")[0] if not sec.startswith("Day") else sec.split(" · ")[0]]=secs.get(sec.split(" · ")[0],0)+unit*qty_value(q,a_vals)
for k,v in secs.items(): print(f"  {k:20s} {v:>12,.0f}")
for k, v in totals.items():
    print(f"{k:9s} raw {v:>12,.0f} XOF  ${v/fx:>9,.0f}")
print(f"Core full  {full(totals['Core']):>12,.0f} XOF  ${full(totals['Core'])/fx:>9,.0f}")
print(f"Core+opt   {full(totals['Core']+totals['Optional']):>12,.0f} XOF  ${full(totals['Core']+totals['Optional'])/fx:>9,.0f}")
